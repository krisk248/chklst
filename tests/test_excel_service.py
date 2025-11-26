"""Tests for Excel service - RED phase"""

import pytest
from datetime import datetime
from pathlib import Path
import tempfile
import shutil
from openpyxl import load_workbook

from backend.services.excel_service import ExcelService


@pytest.fixture
def temp_reports_dir():
    """Create temporary reports directory for testing"""
    temp_dir = tempfile.mkdtemp(prefix="test_excel_")
    yield temp_dir
    # Cleanup
    shutil.rmtree(temp_dir, ignore_errors=True)


@pytest.fixture
def excel_service(temp_reports_dir):
    """Create Excel service with temporary directory"""
    return ExcelService(base_path=temp_reports_dir)


class TestExcelServiceFileManagement:
    """Test file management operations"""

    def test_get_month_folder_creates_directory(self, excel_service, temp_reports_dir):
        """Should create month folder with correct naming"""
        folder = excel_service.get_month_folder(11, 2025)

        assert folder.exists()
        assert folder.name == "Nov_2025"
        assert folder.parent.name == Path(temp_reports_dir).name

    def test_get_month_folder_idempotent(self, excel_service):
        """Should return same folder on multiple calls"""
        folder1 = excel_service.get_month_folder(11, 2025)
        folder2 = excel_service.get_month_folder(11, 2025)

        assert folder1 == folder2
        assert folder1.exists()

    def test_get_project_file_path(self, excel_service, temp_reports_dir):
        """Should generate correct file path for project"""
        file_path = excel_service.get_project_file("BRHUB", 11, 2025)

        assert file_path.name == "BRHUB.xlsx"
        assert file_path.parent.name == "Nov_2025"
        assert str(temp_reports_dir) in str(file_path)

    def test_get_project_file_sanitizes_name(self, excel_service):
        """Should sanitize project names with special characters"""
        file_path = excel_service.get_project_file("My Project / Name", 11, 2025)

        assert file_path.name == "My_Project___Name.xlsx"
        assert "/" not in file_path.name


class TestExcelServiceWriteOperations:
    """Test write operations"""

    def test_save_deployment_creates_file(self, excel_service):
        """Should create new Excel file when saving first deployment"""
        deployment_data = {
            "jira_id": "PAT-001",
            "project_name": "BRHUB",
            "timestamp": datetime(2025, 11, 26, 10, 30),
            "component_name": "BR-HUB",
            "environment": "QA",
            "vcs_url": "https://github.com/example/repo",
            "developer_name": "John Doe",
            "build_server": "192.168.1.149",
            "deploy_server": "192.168.1.60",
            "db_name": "BR_HUB_QA",
            "db_backup_location": "/backup/path",
            "db_script": "script.sql",
            "build_backup": "/backup/build",
            "build_status": True,
            "deploy_status": True,
            "notes": "Test deployment",
            "deployed_by": "Kannan"
        }

        result = excel_service.save_deployment(deployment_data)

        assert result is True
        file_path = excel_service.get_project_file("BRHUB", 11, 2025)
        assert file_path.exists()

    def test_save_deployment_adds_to_deployments_sheet(self, excel_service):
        """Should add deployment data to Deployments sheet"""
        deployment_data = {
            "jira_id": "PAT-001",
            "project_name": "BRHUB",
            "timestamp": datetime(2025, 11, 26, 10, 30),
            "component_name": "BR-HUB",
            "environment": "QA",
            "vcs_url": "https://github.com/example/repo",
            "developer_name": "John Doe",
            "build_server": "192.168.1.149",
            "deploy_server": "192.168.1.60",
            "db_name": "BR_HUB_QA",
            "db_backup_location": "/backup/path",
            "db_script": "script.sql",
            "build_backup": "/backup/build",
            "build_status": True,
            "deploy_status": True,
            "notes": "Test deployment",
            "deployed_by": "Kannan"
        }

        excel_service.save_deployment(deployment_data)

        # Verify file contents
        file_path = excel_service.get_project_file("BRHUB", 11, 2025)
        wb = load_workbook(file_path)
        ws = wb["Deployments"]

        # Should have header + 1 data row
        assert ws.max_row == 2
        assert ws.cell(2, 1).value == "PAT-001"  # JIRA ID
        assert ws.cell(2, 3).value == "BRHUB"    # Project Name

    def test_save_deployment_appends_to_existing_file(self, excel_service):
        """Should append to existing Excel file"""
        # Save first deployment
        deployment1 = {
            "jira_id": "PAT-001",
            "project_name": "BRHUB",
            "timestamp": datetime(2025, 11, 26, 10, 30),
            "component_name": "BR-HUB",
            "environment": "QA",
            "vcs_url": "https://github.com/example/repo",
            "developer_name": "John Doe",
            "build_server": "192.168.1.149",
            "deploy_server": "192.168.1.60",
            "db_name": "BR_HUB_QA",
            "db_backup_location": "/backup/path",
            "db_script": "script.sql",
            "build_backup": "/backup/build",
            "build_status": True,
            "deploy_status": True,
            "notes": "Deployment 1",
            "deployed_by": "Kannan"
        }
        excel_service.save_deployment(deployment1)

        # Save second deployment
        deployment2 = {
            "jira_id": "PAT-002",
            "project_name": "BRHUB",
            "timestamp": datetime(2025, 11, 27, 11, 45),
            "component_name": "BR-HUB",
            "environment": "PROD",
            "vcs_url": "https://github.com/example/repo",
            "developer_name": "Jane Doe",
            "build_server": "192.168.1.149",
            "deploy_server": "192.168.1.60",
            "db_name": "BR_HUB_PROD",
            "db_backup_location": "/backup/path",
            "db_script": "script.sql",
            "build_backup": "/backup/build",
            "build_status": True,
            "deploy_status": True,
            "notes": "Deployment 2",
            "deployed_by": "Kannan"
        }
        excel_service.save_deployment(deployment2)

        # Verify both deployments exist
        file_path = excel_service.get_project_file("BRHUB", 11, 2025)
        wb = load_workbook(file_path)
        ws = wb["Deployments"]

        assert ws.max_row == 3  # Header + 2 deployments
        assert ws.cell(2, 1).value == "PAT-001"
        assert ws.cell(3, 1).value == "PAT-002"

    def test_save_deployment_formats_timestamp(self, excel_service):
        """Should format timestamp in deployment sheet"""
        deployment_data = {
            "jira_id": "PAT-001",
            "project_name": "BRHUB",
            "timestamp": datetime(2025, 11, 26, 14, 30),
            "component_name": "BR-HUB",
            "environment": "QA",
            "vcs_url": "https://github.com/example/repo",
            "developer_name": "John Doe",
            "build_server": "192.168.1.149",
            "deploy_server": "192.168.1.60",
            "db_name": "BR_HUB_QA",
            "db_backup_location": "/backup/path",
            "db_script": "script.sql",
            "build_backup": "/backup/build",
            "build_status": True,
            "deploy_status": True,
            "notes": "Test",
            "deployed_by": "Kannan"
        }

        excel_service.save_deployment(deployment_data)

        file_path = excel_service.get_project_file("BRHUB", 11, 2025)
        wb = load_workbook(file_path)
        ws = wb["Deployments"]

        timestamp_value = ws.cell(2, 2).value
        assert "26" in str(timestamp_value)
        assert "Nov" in str(timestamp_value) or "2025" in str(timestamp_value)

    def test_save_deployment_creates_history_entry(self, excel_service):
        """Should create entry in History sheet"""
        deployment_data = {
            "jira_id": "PAT-001",
            "project_name": "BRHUB",
            "timestamp": datetime(2025, 11, 26, 10, 30),
            "component_name": "BR-HUB",
            "environment": "QA",
            "vcs_url": "https://github.com/example/repo",
            "developer_name": "John Doe",
            "build_server": "192.168.1.149",
            "deploy_server": "192.168.1.60",
            "db_name": "BR_HUB_QA",
            "db_backup_location": "/backup/path",
            "db_script": "script.sql",
            "build_backup": "/backup/build",
            "build_status": True,
            "deploy_status": True,
            "notes": "Test",
            "deployed_by": "Kannan"
        }

        excel_service.save_deployment(deployment_data)

        file_path = excel_service.get_project_file("BRHUB", 11, 2025)
        wb = load_workbook(file_path)
        ws_history = wb["History"]

        # Should have header + 1 history entry
        assert ws_history.max_row >= 2
        assert ws_history.cell(2, 2).value == "CREATE"  # Action


class TestExcelServiceReadOperations:
    """Test read operations"""

    def test_read_deployments_empty_file(self, excel_service):
        """Should return empty list for non-existent file"""
        deployments = excel_service.read_deployments("NonExistent", 11, 2025)

        assert deployments == []

    def test_read_deployments_returns_data(self, excel_service):
        """Should read deployment data from Excel file"""
        # First save a deployment
        deployment_data = {
            "jira_id": "PAT-001",
            "project_name": "BRHUB",
            "timestamp": datetime(2025, 11, 26, 10, 30),
            "component_name": "BR-HUB",
            "environment": "QA",
            "vcs_url": "https://github.com/example/repo",
            "developer_name": "John Doe",
            "build_server": "192.168.1.149",
            "deploy_server": "192.168.1.60",
            "db_name": "BR_HUB_QA",
            "db_backup_location": "/backup/path",
            "db_script": "script.sql",
            "build_backup": "/backup/build",
            "build_status": True,
            "deploy_status": True,
            "notes": "Test",
            "deployed_by": "Kannan"
        }
        excel_service.save_deployment(deployment_data)

        # Now read it back
        deployments = excel_service.read_deployments("BRHUB", 11, 2025)

        assert len(deployments) == 1
        assert deployments[0]["JIRA PATCH ID"] == "PAT-001"
        assert deployments[0]["Project Name"] == "BRHUB"

    def test_read_all_monthly_deployments(self, excel_service):
        """Should read deployments from all project files in a month"""
        # Save deployments to multiple projects
        projects = ["BRHUB", "ADX-SIP", "MBANK"]
        for i, project in enumerate(projects):
            deployment_data = {
                "jira_id": f"PAT-{i+1:03d}",
                "project_name": project,
                "timestamp": datetime(2025, 11, 26, 10, 30),
                "component_name": f"Component-{i+1}",
                "environment": "QA",
                "vcs_url": "https://github.com/example/repo",
                "developer_name": "John Doe",
                "build_server": "192.168.1.149",
                "deploy_server": "192.168.1.60",
                "db_name": "DB_NAME",
                "db_backup_location": "/backup/path",
                "db_script": "script.sql",
                "build_backup": "/backup/build",
                "build_status": True,
                "deploy_status": True,
                "notes": f"Deployment {i+1}",
                "deployed_by": "Kannan"
            }
            excel_service.save_deployment(deployment_data)

        # Read all monthly deployments
        all_deployments = excel_service.read_all_monthly_deployments(11, 2025)

        assert len(all_deployments) == 3
        project_names = [d["Project Name"] for d in all_deployments]
        assert set(project_names) == {"BRHUB", "ADX-SIP", "MBANK"}


class TestExcelServiceFormatting:
    """Test Excel formatting"""

    def test_new_workbook_has_correct_headers(self, excel_service):
        """Should create workbook with correct headers"""
        deployment_data = {
            "jira_id": "PAT-001",
            "project_name": "BRHUB",
            "timestamp": datetime(2025, 11, 26, 10, 30),
            "component_name": "BR-HUB",
            "environment": "QA",
            "vcs_url": "https://github.com/example/repo",
            "developer_name": "John Doe",
            "build_server": "192.168.1.149",
            "deploy_server": "192.168.1.60",
            "db_name": "BR_HUB_QA",
            "db_backup_location": "/backup/path",
            "db_script": "script.sql",
            "build_backup": "/backup/build",
            "build_status": True,
            "deploy_status": True,
            "notes": "Test",
            "deployed_by": "Kannan"
        }
        excel_service.save_deployment(deployment_data)

        file_path = excel_service.get_project_file("BRHUB", 11, 2025)
        wb = load_workbook(file_path)
        ws = wb["Deployments"]

        expected_headers = [
            "JIRA PATCH ID", "Timestamp", "Project Name", "Component Name",
            "Environment", "SVN/GIT URL", "Developer Name", "Build Server",
            "Deploy Server", "Database Name", "DB Backup Location",
            "Database Script", "Previous Build Backup", "Build Status",
            "Deploy Status", "Notes", "Deployed By"
        ]

        actual_headers = [cell.value for cell in ws[1]]
        assert actual_headers == expected_headers

    def test_header_row_is_formatted(self, excel_service):
        """Should format header row with colors and freeze panes"""
        deployment_data = {
            "jira_id": "PAT-001",
            "project_name": "BRHUB",
            "timestamp": datetime(2025, 11, 26, 10, 30),
            "component_name": "BR-HUB",
            "environment": "QA",
            "vcs_url": "https://github.com/example/repo",
            "developer_name": "John Doe",
            "build_server": "192.168.1.149",
            "deploy_server": "192.168.1.60",
            "db_name": "BR_HUB_QA",
            "db_backup_location": "/backup/path",
            "db_script": "script.sql",
            "build_backup": "/backup/build",
            "build_status": True,
            "deploy_status": True,
            "notes": "Test",
            "deployed_by": "Kannan"
        }
        excel_service.save_deployment(deployment_data)

        file_path = excel_service.get_project_file("BRHUB", 11, 2025)
        wb = load_workbook(file_path)
        ws = wb["Deployments"]

        # Check header cell formatting
        header_cell = ws.cell(1, 1)
        assert header_cell.font.bold is True
        assert header_cell.fill.start_color is not None

    def test_history_sheet_has_correct_headers(self, excel_service):
        """Should create History sheet with correct headers"""
        deployment_data = {
            "jira_id": "PAT-001",
            "project_name": "BRHUB",
            "timestamp": datetime(2025, 11, 26, 10, 30),
            "component_name": "BR-HUB",
            "environment": "QA",
            "vcs_url": "https://github.com/example/repo",
            "developer_name": "John Doe",
            "build_server": "192.168.1.149",
            "deploy_server": "192.168.1.60",
            "db_name": "BR_HUB_QA",
            "db_backup_location": "/backup/path",
            "db_script": "script.sql",
            "build_backup": "/backup/build",
            "build_status": True,
            "deploy_status": True,
            "notes": "Test",
            "deployed_by": "Kannan"
        }
        excel_service.save_deployment(deployment_data)

        file_path = excel_service.get_project_file("BRHUB", 11, 2025)
        wb = load_workbook(file_path)
        ws_history = wb["History"]

        expected_headers = [
            "Timestamp", "Action", "Project", "Component",
            "JIRA ID", "User", "Details", "Status"
        ]

        actual_headers = [cell.value for cell in ws_history[1]]
        assert actual_headers == expected_headers
