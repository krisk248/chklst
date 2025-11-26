"""Tests for deployment dual-write to SQLite and Excel - RED phase"""

import pytest
from datetime import datetime
from pathlib import Path
import tempfile
import shutil
from openpyxl import load_workbook

from backend.services.excel_service import ExcelService


@pytest.fixture
def temp_reports_dir():
    """Create temporary reports directory for testing"""
    temp_dir = tempfile.mkdtemp(prefix="test_excel_")
    yield temp_dir
    shutil.rmtree(temp_dir, ignore_errors=True)


@pytest.fixture
def excel_service(temp_reports_dir):
    """Create Excel service with temporary directory"""
    return ExcelService(base_path=temp_reports_dir)


class TestDualWriteDeployment:
    """Test dual-write functionality for deployments"""

    def test_save_deployment_as_dict(self, excel_service):
        """Should save deployment as dictionary to Excel"""
        deployment_data = {
            "jira_id": "PAT-001",
            "project_name": "BRHUB",
            "timestamp": datetime(2025, 11, 26, 10, 30),
            "component_name": "BR-HUB",
            "environment": "QA",
            "vcs_url": "https://github.com/example/repo",
            "developer_name": "John Doe",
            "build_server": "192.168.1.149",
            "deploy_server": "192.168.1.60",
            "db_name": "BR_HUB_QA",
            "db_backup_location": "/backup/path",
            "db_script": "script.sql",
            "build_backup": "/backup/build",
            "build_status": True,
            "deploy_status": True,
            "notes": "Test deployment",
            "deployed_by": "Kannan"
        }

        result = excel_service.save_deployment(deployment_data)

        assert result is True
        file_path = excel_service.get_project_file("BRHUB", 11, 2025)
        assert file_path.exists()

        # Verify data was written
        wb = load_workbook(file_path)
        ws = wb["Deployments"]
        assert ws.cell(2, 1).value == "PAT-001"

    def test_multiple_deployments_same_project_same_month(self, excel_service):
        """Should append multiple deployments to same Excel file"""
        for i in range(1, 4):
            deployment_data = {
                "jira_id": f"PAT-{i:03d}",
                "project_name": "BRHUB",
                "timestamp": datetime(2025, 11, 26 + i, 10, 30),
                "component_name": "BR-HUB",
                "environment": "QA",
                "vcs_url": "https://github.com/example/repo",
                "developer_name": f"Developer {i}",
                "build_server": "192.168.1.149",
                "deploy_server": "192.168.1.60",
                "db_name": "BR_HUB_QA",
                "db_backup_location": "/backup/path",
                "db_script": "script.sql",
                "build_backup": "/backup/build",
                "build_status": True,
                "deploy_status": True,
                "notes": f"Deployment {i}",
                "deployed_by": "Kannan"
            }
            excel_service.save_deployment(deployment_data)

        # Verify all are in Excel
        deployments = excel_service.read_deployments("BRHUB", 11, 2025)
        assert len(deployments) == 3
        jira_ids = [d["JIRA PATCH ID"] for d in deployments]
        assert "PAT-001" in jira_ids
        assert "PAT-002" in jira_ids
        assert "PAT-003" in jira_ids

    def test_deployments_different_projects_same_month(self, excel_service):
        """Should create separate Excel files for different projects"""
        projects = ["BRHUB", "ADX-SIP", "MBANK"]
        for i, project in enumerate(projects):
            deployment_data = {
                "jira_id": f"PAT-{i+1:03d}",
                "project_name": project,
                "timestamp": datetime(2025, 11, 26, 10, 30),
                "component_name": f"Component-{i+1}",
                "environment": "QA",
                "vcs_url": "https://github.com/example/repo",
                "developer_name": "John Doe",
                "build_server": "192.168.1.149",
                "deploy_server": "192.168.1.60",
                "db_name": "DB_NAME",
                "db_backup_location": "/backup/path",
                "db_script": "script.sql",
                "build_backup": "/backup/build",
                "build_status": True,
                "deploy_status": True,
                "notes": f"Deployment for {project}",
                "deployed_by": "Kannan"
            }
            excel_service.save_deployment(deployment_data)

        # Verify files were created
        file_brhub = excel_service.get_project_file("BRHUB", 11, 2025)
        file_adx = excel_service.get_project_file("ADX-SIP", 11, 2025)
        file_mbank = excel_service.get_project_file("MBANK", 11, 2025)

        assert file_brhub.exists()
        assert file_adx.exists()
        assert file_mbank.exists()

        # Verify each file has one deployment
        assert len(excel_service.read_deployments("BRHUB", 11, 2025)) == 1
        assert len(excel_service.read_deployments("ADX-SIP", 11, 2025)) == 1
        assert len(excel_service.read_deployments("MBANK", 11, 2025)) == 1

    def test_deployments_different_months(self, excel_service):
        """Should create separate files for different months"""
        # Deployment in November
        deployment_nov = {
            "jira_id": "PAT-001",
            "project_name": "BRHUB",
            "timestamp": datetime(2025, 11, 26, 10, 30),
            "component_name": "BR-HUB",
            "environment": "QA",
            "vcs_url": "https://github.com/example/repo",
            "developer_name": "John Doe",
            "build_server": "192.168.1.149",
            "deploy_server": "192.168.1.60",
            "db_name": "BR_HUB_QA",
            "db_backup_location": "/backup/path",
            "db_script": "script.sql",
            "build_backup": "/backup/build",
            "build_status": True,
            "deploy_status": True,
            "notes": "November deployment",
            "deployed_by": "Kannan"
        }
        excel_service.save_deployment(deployment_nov)

        # Deployment in December
        deployment_dec = {
            "jira_id": "PAT-002",
            "project_name": "BRHUB",
            "timestamp": datetime(2025, 12, 1, 10, 30),
            "component_name": "BR-HUB",
            "environment": "QA",
            "vcs_url": "https://github.com/example/repo",
            "developer_name": "Jane Doe",
            "build_server": "192.168.1.149",
            "deploy_server": "192.168.1.60",
            "db_name": "BR_HUB_QA",
            "db_backup_location": "/backup/path",
            "db_script": "script.sql",
            "build_backup": "/backup/build",
            "build_status": True,
            "deploy_status": True,
            "notes": "December deployment",
            "deployed_by": "Kannan"
        }
        excel_service.save_deployment(deployment_dec)

        # Verify files are in different folders
        nov_deployments = excel_service.read_deployments("BRHUB", 11, 2025)
        dec_deployments = excel_service.read_deployments("BRHUB", 12, 2025)

        assert len(nov_deployments) == 1
        assert nov_deployments[0]["JIRA PATCH ID"] == "PAT-001"

        assert len(dec_deployments) == 1
        assert dec_deployments[0]["JIRA PATCH ID"] == "PAT-002"
